"""Studi pola IHSG: apakah kumpulan sinyal & rezim EMA200 mendahului rally / penurunan?

Untuk setiap hari bursa, dihitung (as-of penutupan hari itu, tanpa melihat masa depan):
  - jumlah candle bersinyal SELL / BUY dalam N hari bursa terakhir (Daily dan/atau Weekly;
    candle Weekly dihitung pada hari terakhir minggunya). Meter STRONG memakai SMA50.
  - rezim EMA200 (UPTREND / DOWNTREND / TRANSISI)
lalu diukur apa yang terjadi pada IHSG 5, 10, dan 20 hari bursa berikutnya.

Hasil: Excel (ringkasan aturan, sinyal per rezim, daftar episode) + ringkasan Telegram.
"""
import logging

import numpy as np
import pandas as pd

from . import data, indicators as ta
from .market import SLOPE_BARS, closed_weeks, regime, weekly_bars
from .signals import candles, evaluate

log = logging.getLogger("macroflow.study")
ALL = ("COMBO", "STANDAR", "METER")
HORIZONS = (5, 10, 20)


def _signals(chart, daily, tf, p):
    """Semua sinyal indeks. Volume di-nolkan supaya ROYAL (mati untuk indeks) tidak dihitung."""
    c = chart.copy()
    c["Volume"] = 0.0
    d = daily.copy()
    d["Volume"] = 0.0
    sigs = evaluate(c, d, None, tf, p, last_n=len(c) - 1)
    return pd.DataFrame([{"time": s.bar_time, "tf": tf, "name": s.name, "side": s.side}
                         for s in sigs if s.name != "ROYAL"])


def _day_counts(sig, days, side, names, tfs):
    """Jumlah candle bersinyal per hari (beberapa sinyal di candle yang sama = 1)."""
    if sig.empty:
        return pd.Series(0, index=days)
    s = sig[(sig.side == side) & sig.name.isin(names) & sig.tf.isin(tfs)]
    s = s.drop_duplicates(["time", "tf"])
    cnt = s.groupby(s.time.dt.normalize()).size()
    return cnt.reindex(days, fill_value=0)


def build_daily(d, p, window=5, now=None):
    """Tabel harian: close, rezim, return ke depan, dan hitungan sinyal untuk berbagai aturan."""
    days = d.index
    c = d["Close"].to_numpy(float)       # harga asli: untuk return ke depan
    lo = d["Low"].to_numpy(float)
    hi = d["High"].to_numpy(float)
    hc = candles(d, p)["Close"].to_numpy(float)  # rezim dari candle yang dipakai indikator (HA)
    e200 = ta.ema(hc, 200)
    prev = np.concatenate([np.full(SLOPE_BARS, np.nan), e200[:-SLOPE_BARS]])
    reg = [regime(hc[i], e200[i], prev[i]) if not np.isnan(prev[i]) else None for i in range(len(c))]

    sig = _signals(d, d, "D", p)
    now = d.index[-1] + pd.Timedelta(days=1) if now is None else now
    wk = closed_weeks(weekly_bars(d), d.index[-1], now)
    if wk is not None and len(wk) >= 30:
        sig = pd.concat([sig, _signals(wk, d, "W", p)], ignore_index=True)

    df = pd.DataFrame({"Close": c, "Rezim": reg}, index=days)
    df["EMA200"] = e200
    n = len(c)
    for k in HORIZONS:
        fut = np.full(n, np.nan)
        fut[:n - k] = c[k:] / c[:n - k] * 100 - 100
        df[f"Ret {k}h %"] = fut
    dd = np.full(n, np.nan)
    up = np.full(n, np.nan)
    for i in range(n - 1):
        j = min(n, i + 21)
        dd[i] = lo[i + 1:j].min() / c[i] * 100 - 100
        up[i] = hi[i + 1:j].max() / c[i] * 100 - 100
    df["Max turun 20h %"] = dd
    df["Max naik 20h %"] = up

    for side in ("SELL", "BUY"):
        for label, names, tfs in (("semua", ALL, ("D", "W")), ("semuaD", ALL, ("D",)),
                                  ("strong", ("METER",), ("D", "W"))):
            per_day = _day_counts(sig, days, side, names, tfs)
            df[f"{side}_{label}"] = per_day.rolling(window, min_periods=1).sum().astype(int).values
    return df, sig


def rules(df, window=5):
    """(nama aturan, arah yang diharapkan, mask hari aktif)"""
    R = []
    dn, upt = df["Rezim"] == "DOWNTREND", df["Rezim"] == "UPTREND"
    R.append((f"STRONG SELL ≥2× in {window} days (D+W)", "turun", df.SELL_strong >= 2))
    R.append((f"Any SELL on ≥2 candles in {window} days (D+W) — active rule", "turun", df.SELL_semua >= 2))
    R.append((f"Any SELL on ≥3 candles in {window} days (D+W)", "turun", df.SELL_semua >= 3))
    R.append((f"Any SELL on ≥2 candles in {window} days (Daily only)", "turun", df.SELL_semuaD >= 2))
    R.append(("Active rule + DOWNTREND", "turun", (df.SELL_semua >= 2) & dn))
    R.append(("DOWNTREND (regime only)", "turun", dn))
    R.append(("UPTREND (regime only)", "naik", upt))
    R.append((f"Any BUY on ≥2 candles in {window} days (D+W)", "naik", df.BUY_semua >= 2))
    R.append(("BUY ≥2 (D+W) + UPTREND — uptrend confirmation", "naik", (df.BUY_semua >= 2) & upt))
    return R


def _episodes(mask, gap=5):
    """Hari pertama aturan aktif setelah minimal `gap` hari tidak aktif."""
    m = mask.to_numpy()
    out = np.zeros(len(m), bool)
    last = -10 ** 9
    for i, v in enumerate(m):
        if v:
            if i - last > gap:
                out[i] = True
            last = i
    return pd.Series(out, index=mask.index)


def summarize(df, window=5):
    rows, eps = [], []
    for name, expect, mask in rules(df, window):
        base = df
        mask = mask.fillna(False).astype(bool)
        ep = _episodes(mask)
        e = base[ep]
        good = (e["Ret 10h %"] < 0) if expect == "turun" else (e["Ret 10h %"] > 0)
        r = {"Aturan": name, "Arah diharapkan": expect,
             "Periode": f"{base.index[0]:%b %Y} – {base.index[-1]:%b %Y}" if len(base) else "-",
             "Hari aktif": int(mask.sum()), "% hari": mask.mean() * 100 if len(base) else np.nan,
             "Episode": int(ep.sum()),
             "Arah benar 10h %": good[e["Ret 10h %"].notna()].mean() * 100 if len(e) else np.nan}
        for k in HORIZONS:
            r[f"Rata2 ret {k}h % (episode)"] = e[f"Ret {k}h %"].mean()
            r[f"Rata2 ret {k}h % (semua hari)"] = base[f"Ret {k}h %"].mean()
        r["Rata2 max turun 20h %"] = e["Max turun 20h %"].mean()
        r["Turun ≥5% dalam 20h (%)"] = (e["Max turun 20h %"] <= -5).mean() * 100 if len(e) else np.nan
        r["Baseline turun ≥5% dalam 20h (%)"] = (base["Max turun 20h %"] <= -5).mean() * 100
        rows.append(r)
        for t, x in e.iterrows():
            eps.append({"Aturan": name, "Tanggal": t, "Close": x.Close, "Rezim": x.Rezim,
                        **{f"Ret {k}h %": x[f"Ret {k}h %"] for k in HORIZONS},
                        "Max turun 20h %": x["Max turun 20h %"], "Max naik 20h %": x["Max naik 20h %"]})
    return pd.DataFrame(rows), pd.DataFrame(eps)


def by_regime(df, sig):
    """Kinerja sinyal IHSG (Daily & Weekly) dikelompokkan menurut rezim saat sinyal muncul."""
    s = sig.drop_duplicates(["time", "tf", "side"]).copy()
    if s.empty:
        return pd.DataFrame()
    s["Rezim"] = s.time.map(df["Rezim"])
    s["Ret 10h %"] = s.time.map(df["Ret 10h %"])
    s["Ret 20h %"] = s.time.map(df["Ret 20h %"])
    s = s.dropna(subset=["Rezim", "Ret 10h %"])
    s["benar"] = np.where(s.side == "BUY", s["Ret 10h %"] > 0, s["Ret 10h %"] < 0)
    s["TF"] = s.tf
    g = s.groupby(["TF", "side", "Rezim"]).agg(
        Jumlah=("benar", "size"), **{"Arah benar 10h %": ("benar", lambda x: x.mean() * 100)},
        **{"Rata2 ret 10h %": ("Ret 10h %", "mean")}, **{"Rata2 ret 20h %": ("Ret 20h %", "mean")})
    return g.reset_index().rename(columns={"side": "Arah sinyal"})


def write_excel(out, summary, regime_tbl, eps, df):
    from openpyxl.styles import Alignment, Font, PatternFill

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Aturan", index=False)
        regime_tbl.to_excel(xw, sheet_name="Sinyal per rezim", index=False)
        e = eps.copy()
        if len(e):
            e["Tanggal"] = e["Tanggal"].dt.strftime("%Y-%m-%d")
        e.to_excel(xw, sheet_name="Episode", index=False)
        h = df.rename_axis("Tanggal").reset_index()  # index Yahoo bernama "Date"
        h["Tanggal"] = pd.to_datetime(h["Tanggal"]).dt.strftime("%Y-%m-%d")
        h.to_excel(xw, sheet_name="Harian", index=False)
        pd.DataFrame({"Catatan": [
            "Setiap hari dihitung as-of penutupan: jumlah candle bersinyal dalam 5 hari bursa terakhir.",
            "Beberapa sinyal di candle yang sama dihitung 1. 'Semua' = COMBO, STANDAR, METER (STRONG).",
            "Episode = hari pertama aturan aktif setelah minimal 5 hari bursa tidak aktif.",
            "Weekly: candle mingguan yang sudah selesai, dihitung pada hari bursa terakhir minggunya.",
            "Meter STRONG memakai SMA50 (sama dengan alert indeks).",
            "Rezim: UPTREND = harga > EMA200 dan EMA200 naik 20 hari; DOWNTREND = kebalikannya.",
            "Bandingkan kolom (episode) dengan (semua hari): selisih besar = aturan punya informasi.",
            "Sampel kecil (puluhan episode) — anggap sebagai petunjuk, bukan kepastian.",
        ]}).to_excel(xw, sheet_name="Catatan", index=False)
        fill, font = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.fill, cell.font = fill, font
                cell.alignment = Alignment(wrap_text=True, horizontal="center")
            for col in ws.columns:
                letter = col[0].column_letter
                ws.column_dimensions[letter].width = 95 if ws.title == "Catatan" else (
                    48 if col[0].value == "Aturan" else 14)
                for cell in col[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0"
            ws.freeze_panes = "B2"


def telegram_text(summary, regime_tbl):
    REG = {"TRANSISI": "TRANSITION"}
    lines = ["🔬 <b>IHSG PATTERN STUDY</b>",
             "<i>Episode = first day a rule becomes active. Compare with the average of all days.</i>"]
    for _, r in summary.iterrows():
        k = "Rata2 ret 10h % (episode)"
        b = "Rata2 ret 10h % (semua hari)"
        if not r["Episode"]:
            lines.append(f"\n<b>{r['Aturan']}</b>: never active")
            continue
        lines.append(
            f"\n<b>{r['Aturan']}</b>\n{r['Episode']} episodes · right direction after 10d "
            f"{r['Arah benar 10h %']:.0f}% · 10d return {r[k]:+.1f}% (normal {r[b]:+.1f}%) · "
            f"fell ≥5% within 20d: {r['Turun ≥5% dalam 20h (%)']:.0f}% "
            f"(normal {r['Baseline turun ≥5% dalam 20h (%)']:.0f}%)")
    if len(regime_tbl):
        lines.append("\n<b>IHSG signals by regime</b> (right direction after 10 trading days)")
        for _, r in regime_tbl.iterrows():
            lines.append(f"• {r['TF']} {r['Arah sinyal']} in {REG.get(r['Rezim'], r['Rezim'])}: "
                         f"{r['Jumlah']}× · right {r['Arah benar 10h %']:.0f}% · "
                         f"10d return {r['Rata2 ret 10h %']:+.1f}%")
    lines.append("\nDetails: Excel file <b>ihsg-study</b> on the GitHub run page.")
    return "\n".join(lines)


def run(cfg, p, out, code="IHSG", window=5, daily_period="10y", now=None):
    from .index_alert import index_params

    mw = cfg.get("market_warning") or {}
    code = str(mw.get("index", code)).upper()
    window = int(mw.get("window_days", window))
    p = index_params(cfg, p)  # Meter STRONG dengan SMA, sama dengan alert indeks
    d = data.download([code], daily_period, "1d").get(code)
    if d is None or len(d) < 260:
        raise RuntimeError(f"Data harian {code} tidak tersedia")
    if now is not None:  # buang bar hari ini kalau belum close
        from .index_alert import final_upto
        fu = final_upto(d, now, "IDX")
        d = d[d.index <= fu] if fu is not None else d
    log.info("Studi %s: %d hari (Daily & Weekly, Meter %s)", code, len(d), p.meter_ma)
    df, sig = build_daily(d, p, window, now)
    summary, eps = summarize(df, window)
    reg = by_regime(df, sig)
    write_excel(out, summary, reg, eps, df)
    return summary, reg
