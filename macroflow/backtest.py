"""Backtest: putar ulang aplikasi pada periode lalu & ukur hasil tiap sinyal.

Asumsi:
- Sinyal dihitung persis seperti mode live (Daily = candle final, 1H = candle yang sudah close).
- Entry = harga close candle sinyal.
- BUY: dicek candle-candle berikutnya (timeframe yang sama) mana yang tersentuh duluan,
  SL atau TP1/TP2. Kalau SL & TP tersentuh di candle yang sama -> dianggap SL (konservatif).
- SELL: diukur dari perubahan harga setelah 5 & 10 candle (turun = sinyal benar).
- Universe = daftar saham hari ini, filter likuiditas dihitung per tanggal mulai.
"""
import logging
from pathlib import Path

import numpy as np
import pandas as pd

from . import data, notify
from .signals import evaluate

log = logging.getLogger("macroflow.backtest")


def simulate(chart, i, sig):
    """Hasil satu sinyal. chart = OHLCV timeframe sinyal, i = posisi bar sinyal."""
    c = chart["Close"].to_numpy(float)
    hi = chart["High"].to_numpy(float)
    lo = chart["Low"].to_numpy(float)
    entry = c[i]
    n = len(c)

    def ret(k):
        return (c[i + k] / entry - 1) * 100 if i + k < n else np.nan

    row = {
        "Ret 5 bar %": ret(5), "Ret 10 bar %": ret(10),
        "Ret s/d kini %": (c[-1] / entry - 1) * 100,
    }
    fut_hi, fut_lo = hi[i + 1:], lo[i + 1:]
    row["Max naik %"] = (fut_hi.max() / entry - 1) * 100 if len(fut_hi) else np.nan
    row["Max turun %"] = (fut_lo.min() / entry - 1) * 100 if len(fut_lo) else np.nan

    if sig.side == "BUY" and "sl" in sig.info:
        sl, tp1, tp2 = sig.info["sl"], sig.info["tp1"], sig.info["tp2"]
        result, bars = "Berjalan", np.nan
        hit_tp1 = False
        for j in range(i + 1, n):
            if lo[j] <= sl:
                result = "TP1" if hit_tp1 else "SL"
                bars = j - i
                break
            if hi[j] >= tp2:
                result, bars = "TP2", j - i
                break
            if hi[j] >= tp1 and not hit_tp1:
                hit_tp1, bars = True, j - i
        else:
            if hit_tp1:
                result = "TP1"
        row.update({"SL": sl, "TP1": tp1, "TP2": tp2, "Hasil": result, "Bar ke hasil": bars})
    elif sig.side == "SELL":
        r5 = row["Ret 5 bar %"]
        row["Hasil"] = "-" if np.isnan(r5) else ("Benar (turun)" if r5 < 0 else "Salah (naik)")
    return row


def run(cfg, start, end=None, tfs=("D", "1H"), out=Path("backtest_hasil.xlsx"),
        index_sigs=None, enabled=None, allowed=None, params=None, liquid=None, root=Path("."),
        now=None):
    start = pd.Timestamp(start)
    end = pd.Timestamp(end) if end else None
    index_sigs = index_sigs or {}

    tickers = data.load_universe(cfg, root / "tickers.txt")
    daily = data.download(tickers, cfg.get("daily_period", "3y"), "1d")
    flt = cfg.get("filter") or {}
    keep = [t for t, df in daily.items()
            if liquid(df, float(flt.get("min_avg_value", 2e9)), float(flt.get("min_price", 50)), start)]
    idx = [str(c).upper() for c in (cfg.get("indices") or {})]  # indeks ikut scan (default: tidak ada)
    if idx:
        daily.update(data.download(idx, cfg.get("daily_period", "3y"), "1d"))
        keep = idx + [t for t in keep if t not in idx]
    log.info("Backtest %d saham (lolos filter per %s)", len(keep), start.date())
    hourly = data.download(keep, cfg.get("hourly_period", "1y"), "1h")

    rows = []
    for k, t in enumerate(keep, 1):
        d, h = daily.get(t), hourly.get(t)
        if d is None or h is None:
            continue
        for tf in tfs:
            chart = d if tf == "D" else h
            if now is not None:  # buang candle yang belum close
                if tf == "1H":
                    chart = chart[chart.index + pd.Timedelta(hours=1) <= now]
                elif now.hour < 16:
                    chart = chart[chart.index < now.normalize()]
            sel = chart.index >= start
            if end is not None:
                sel &= chart.index < end + pd.Timedelta(days=1)
            n_bt = int(sel.sum())
            if n_bt == 0:
                continue
            allow = enabled(cfg, tf)
            try:
                last_i = int(np.where(sel)[0][-1])
                sigs = evaluate(chart.iloc[:last_i + 1], d, h, tf, params, last_n=n_bt)
            except Exception as e:
                log.warning("%s %s gagal: %s", t, tf, e)
                continue
            pos = {ts: i for i, ts in enumerate(chart.index)}
            for s in sigs:
                if not allowed(s, t, allow, index_sigs) or s.bar_time < start:
                    continue
                r = {"Waktu": s.bar_time, "Saham": t, "TF": tf, "Sinyal": s.name, "Arah": s.side,
                     "Close": s.close, "RSI": s.info.get("rsi"), "Vol x": s.info.get("vol_ratio"),
                     "Lapis SMC": s.info.get("smc_level"),
                     "Zona SMC": ("{:.2f}–{:.2f}".format(*s.info["zone"]) if s.info.get("zone") else None)}
                r.update(simulate(chart, pos[s.bar_time], s))
                rows.append(r)
        if k % 50 == 0:
            log.info("... %d/%d saham", k, len(keep))

    df = pd.DataFrame(rows)
    if df.empty:
        log.info("Tidak ada sinyal pada periode ini")
        return df, pd.DataFrame()
    df = df.sort_values(["TF", "Waktu", "Saham"]).reset_index(drop=True)
    cols = ["Waktu", "Saham", "TF", "Sinyal", "Arah", "Close", "SL", "TP1", "TP2", "Hasil",
            "Bar ke hasil", "Ret 5 bar %", "Ret 10 bar %", "Ret s/d kini %", "Max naik %",
            "Max turun %", "RSI", "Vol x", "Lapis SMC", "Zona SMC"]
    df = df[[c for c in cols if c in df] + [c for c in df if c not in cols]]
    summary = summarize(df)
    write_excel(df, summary, out, start, end, len(keep))
    log.info("Hasil backtest: %s (%d sinyal)", out, len(df))
    return df, summary


def summarize(df):
    out = []
    for (tf, name, side), g in df.groupby(["TF", "Sinyal", "Arah"], sort=False):
        r5 = g["Ret 5 bar %"]
        row = {"TF": tf, "Sinyal": name, "Arah": side, "Jumlah": len(g),
               "Rata2 ret 5 bar %": r5.mean(), "Rata2 ret 10 bar %": g["Ret 10 bar %"].mean()}
        good = (r5 > 0) if side == "BUY" else (r5 < 0)
        row["Arah benar 5 bar %"] = good[r5.notna()].mean() * 100 if r5.notna().any() else np.nan
        if side == "BUY" and "Hasil" in g:
            res = g["Hasil"]
            closed = res.isin(["SL", "TP1", "TP2"]).sum()
            row.update({
                "TP2": int((res == "TP2").sum()), "TP1": int((res == "TP1").sum()),
                "SL": int((res == "SL").sum()), "Berjalan": int((res == "Berjalan").sum()),
                "Win rate %": (res.isin(["TP1", "TP2"]).sum() / closed * 100) if closed else np.nan,
            })
        out.append(row)
    order = {n: i for i, n in enumerate(notify.ORDER)}
    s = pd.DataFrame(out)
    for col in ("TP2", "TP1", "SL", "Berjalan"):
        if col in s:
            s[col] = s[col].astype("Int64")
    s["_o"] = s["TF"].map({"D": 0, "1H": 1}) * 100 + s["Sinyal"].map(order) * 2 + (s["Arah"] == "SELL")
    return s.sort_values("_o").drop(columns="_o").reset_index(drop=True)


def write_excel(df, summary, out, start, end, n_stocks):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    fills = {"TP2": "C6EFCE", "TP1": "E2F0D9", "SL": "F8CBAD", "Benar (turun)": "C6EFCE",
             "Salah (naik)": "F8CBAD"}
    period = f"{start:%d %b %Y} – {(end or pd.Timestamp.today()):%d %b %Y}"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Ringkasan", index=False, startrow=3)
        for tf, label in (("D", "Sinyal Daily"), ("1H", "Sinyal 1H")):
            part = df[df["TF"] == tf].drop(columns="TF")
            if len(part):
                part = part.copy()
                fmt = "%Y-%m-%d" if tf == "D" else "%Y-%m-%d %H:%M"
                part["Waktu"] = part["Waktu"].dt.strftime(fmt)
                part.dropna(axis=1, how="all").to_excel(xw, sheet_name=label, index=False)
        info = pd.DataFrame({"Catatan": [
            f"Periode: {period}. Saham dicek: {n_stocks} (lolos filter likuiditas per tanggal mulai).",
            "Entry = close candle sinyal. Sinyal 1H = candle yang sudah close (seperti mode live).",
            "BUY: hasil = level yang tersentuh duluan di candle berikutnya. SL & TP di candle sama = SL.",
            "TP1 = sempat menyentuh TP1 sebelum SL (belum/tidak sampai TP2). Berjalan = belum menyentuh apa pun.",
            "Win rate = (TP1 + TP2) / (TP1 + TP2 + SL), sinyal 'Berjalan' tidak dihitung.",
            "SELL: benar bila harga 5 candle kemudian lebih rendah dari close sinyal.",
            "Data: Yahoo Finance. Tidak termasuk biaya transaksi, slippage, ARB/ARA, dan antrian.",
            "Hasil masa lalu tidak menjamin hasil di masa depan.",
        ]})
        info.to_excel(xw, sheet_name="Catatan", index=False)

        wb = xw.book
        ws = wb["Ringkasan"]
        ws["A1"] = "Backtest MacroFlow V65"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = period
        for ws in wb.worksheets:
            hdr = 4 if ws.title == "Ringkasan" else 1
            for cell in ws[hdr]:
                cell.fill, cell.font = head_fill, head_font
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
            names = [c.value for c in ws[hdr]]
            for col, name in enumerate(names, 1):
                letter = get_column_letter(col)
                width = 90 if ws.title == "Catatan" else max(9, min(18, len(str(name or "")) + 2))
                ws.column_dimensions[letter].width = width
                for cell in ws[letter][hdr:]:
                    if isinstance(cell.value, float):
                        if name in ("Close", "SL", "TP1", "TP2"):
                            cell.number_format = "#,##0"
                        elif name in ("RSI", "Bar ke hasil"):
                            cell.number_format = "0"
                        else:
                            cell.number_format = "0.0"
                    if name in ("Hasil",) and cell.value in fills:
                        cell.fill = PatternFill("solid", fgColor=fills[cell.value])
                    if name == "Arah":
                        cell.font = Font(bold=True, color="00B050" if cell.value == "BUY" else "C00000")
            ws.auto_filter.ref = ws.dimensions if ws.title != "Ringkasan" else \
                f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"


def telegram_summary(summary, start, end):
    lines = [f"📊 <b>BACKTEST MacroFlow</b>\n{start:%d %b %Y} – {(end or pd.Timestamp.today()):%d %b %Y}"]
    for tf, g in summary.groupby("TF", sort=False):
        lines.append(f"\n<b>Timeframe {tf}</b>")
        for _, r in g.iterrows():
            emoji = "🟢" if r["Arah"] == "BUY" else "🔴"
            name = "STANDARD" if r["Sinyal"] == "STANDAR" else r["Sinyal"]
            txt = f"{emoji} {name} {r['Arah']}: {r['Jumlah']} signals"
            if r["Arah"] == "BUY" and not pd.isna(r.get("Win rate %")):
                txt += f" | win {r['Win rate %']:.0f}% (TP2 {r['TP2']}, TP1 {r['TP1']}, SL {r['SL']})"
            if not pd.isna(r["Arah benar 5 bar %"]):
                txt += f" | right direction after 5 bars {r['Arah benar 5 bar %']:.0f}%"
            lines.append(txt)
    lines.append("\nFull details are in the backtest Excel file.")
    return "\n".join(lines)
